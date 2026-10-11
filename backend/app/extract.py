"""Turn downloaded file bytes into plain text for indexing."""

import io

from docx import Document
from pypdf import PdfReader


def _txt(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(part for part in parts if part)


_EXTRACTORS = {
    "txt": _txt,
    "gdoc": _txt,  # Google Docs are exported as plain text
    "pdf": _pdf,
    "docx": _docx,
}


def extract_text(file_type: str, data: bytes) -> str:
    return _EXTRACTORS[file_type](data)
