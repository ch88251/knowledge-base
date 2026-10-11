import io

from docx import Document

from app.extract import extract_text


def test_txt_handles_bom_and_non_utf8():
    assert extract_text("txt", "﻿hello".encode("utf-8")) == "hello"
    assert extract_text("txt", "café".encode("latin-1")) == "café"


def test_docx_includes_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph("Resetting a password")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "table cell text"
    buffer = io.BytesIO()
    doc.save(buffer)

    text = extract_text("docx", buffer.getvalue())
    assert "Resetting a password" in text
    assert "table cell text" in text


def test_pdf():
    # Minimal one-page PDF containing the text "Hello KB".
    stream = b"BT /F1 24 Tf 72 720 Td (Hello KB) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(pdf)
    pdf += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        pdf += b"%010d 00000 n \n" % offset
    pdf += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, xref)

    assert "Hello KB" in extract_text("pdf", pdf)
